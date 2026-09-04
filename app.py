import io
import json
import os
import re
import xml.etree.ElementTree as ET
import zipfile
from datetime import datetime, timezone
from functools import wraps
from bson.objectid import ObjectId
from flask import Flask, jsonify, redirect, render_template, request, session, url_for
from pymongo import MongoClient
from werkzeug.security import check_password_hash, generate_password_hash

app = Flask(__name__)
app.secret_key = os.urandom(32).hex()

# =========================================================================
# 🌐 الاتصال بقاعدة بيانات MongoDB Atlas السحابية
# =========================================================================
MONGO_URI = "mongodb+srv://ahmedosman:i-fn%40bBHV7rXMYj@cluster0.8wawfsu.mongodb.net/?retryWrites=true&w=majority&appName=Cluster0"

try:
    client = MongoClient(MONGO_URI, maxPoolSize=50, connectTimeoutMS=5000, socketTimeoutMS=5000)
    db = client['haidy_edu_platform']
    users_col = db['users']
    attendance_col = db['attendance']
    sessions_col = db['sessions']
    exams_col = db['exams']
    results_col = db['results']
    print("✅ تم الاتصال بقاعدة بيانات MongoDB Atlas بنجاح!")
except Exception as e:
    print(f"❌ خطأ في الاتصال بقاعدة البيانات: {e}")

GRADE_NAMES = {
    'p4': 'الصف الرابع الابتدائي',
    'p5': 'الصف الخامس الابتدائي',
    'p6': 'الصف السادس الابتدائي',
    'm1': 'الصف الأول الإعدادي',
    'm2': 'الصف الثاني الإعدادي'
}

DEFAULT_AVATAR = "https://images.unsplash.com/photo-1573496359142-b8d87734a5a2?q=80&w=600&auto=format&fit=crop"

def init_admin():
    try:
        admin = users_col.find_one({"role": "admin"})
        if not admin:
            users_col.insert_one({
                "name": "الأستاذة هايدي عطية",
                "phone": "01000000000",
                "password": generate_password_hash("Haidy@2026"),
                "role": "admin",
                "student_code": "TEACHER-HAIDY",
                "grade": "all",
                "status": "approved",
                "avatar": DEFAULT_AVATAR,
                "created_at": datetime.now(timezone.utc)
            })
    except Exception as e:
        print(f"Admin init warning: {e}")

init_admin()

def login_required(role=None):
    def decorator(f):
        @wraps(f)
        def decorated_function(*args, **kwargs):
            if 'user_id' not in session:
                return redirect(url_for('index'))
            if role and session.get('role') != role:
                return redirect(url_for('index'))
            return f(*args, **kwargs)
        return decorated_function
    return decorator

# =========================================================================
# 🛠️ محرك قراءة واستخراج الأسئلة من Word
# =========================================================================
def clean_hidden_chars(text):
    bad_chars = [
        '\ufeff', '\xa0', '\u200e', '\u200f', '\u202a', '\u202b', 
        '\u202c', '\u202d', '\u202e', '\u200b', '\u200c', '\u200d'
    ]
    for ch in bad_chars:
        text = text.replace(ch, ' ')
    return text

def extract_text_from_docx(file_storage):
    file_bytes = file_storage.read()
    paragraphs = []
    try:
        from docx import Document
        doc = Document(io.BytesIO(file_bytes))
        for p in doc.paragraphs:
            cleaned = clean_hidden_chars(p.text).strip()
            if cleaned:
                paragraphs.append(cleaned)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        cleaned = clean_hidden_chars(p.text).strip()
                        if cleaned and cleaned not in paragraphs:
                            paragraphs.append(cleaned)
    except Exception:
        try:
            with zipfile.ZipFile(io.BytesIO(file_bytes)) as z:
                xml_content = z.read('word/document.xml')
            root = ET.fromstring(xml_content)
            ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
            for p in root.findall('.//w:p', ns):
                t = "".join([node.text for node in p.findall('.//w:t', ns) if node.text])
                cleaned = clean_hidden_chars(t).strip()
                if cleaned:
                    paragraphs.append(cleaned)
        except Exception as ex:
            print("Docx Fallback Error:", ex)

    return "\n".join(paragraphs)

def parse_mcq_content(text):
    text = clean_hidden_chars(text)
    text = re.sub(r'([^\s\n])\s*([A-Dأبجد][\.:\)])', r'\1\n\2', text)
    lines = [l.strip() for l in text.splitlines() if l.strip()]
    
    questions = []
    current_q = None
    
    q_pattern = re.compile(r'^(\d+)[\.:\)\-]\s*(.+)')
    opt_pattern = re.compile(r'^[\(\[\{]?\s*([A-Da-dأبجد])\s*[\.\:\)\-\]\=\s]\s*(.+)')

    arabic_map = {
        'أ': 'a', 'ا': 'a', 'إ': 'a', 'آ': 'a',
        'ب': 'b',
        'ج': 'c',
        'د': 'd'
    }

    for line in lines:
        cleaned_line = clean_hidden_chars(line).strip()
        if not cleaned_line:
            continue

        if re.search(r'(?:[إا]ل?[إا]?جاب[ةه]|جواب|حل|correct|answer|ans)', cleaned_line, re.IGNORECASE):
            if current_q:
                clean_ans = re.sub(
                    r'^(?:.*?[إا]ل?[إا]?جاب[ةه].*?|.*?correct.*?|.*?answer.*?|.*?حل.*?|.*?جواب.*?)(?:[\s:=ـ\-]|\s+هي\s+)+',
                    '',
                    cleaned_line,
                    flags=re.IGNORECASE
                ).strip()
                
                match = re.search(r'([A-Da-dأبجد])', clean_ans)
                if match:
                    raw_char = match.group(1).lower()
                    current_q["correct"] = arabic_map.get(raw_char, raw_char)
                else:
                    ans_lower = clean_ans.lower()
                    for idx, opt in enumerate(current_q.get("options", [])):
                        if opt.lower() in ans_lower or ans_lower in opt.lower():
                            current_q["correct"] = ['a', 'b', 'c', 'd'][idx]
                            break
            continue

        q_match = q_pattern.match(cleaned_line)
        if q_match:
            if current_q:
                current_q["type"] = "tf" if len(current_q["options"]) == 2 else "mcq"
                questions.append(current_q)
            current_q = {
                "id": len(questions) + 1,
                "text": q_match.group(2).strip(),
                "options": [],
                "correct": "a",
                "type": "mcq"
            }
            continue

        if current_q:
            opt_match = opt_pattern.match(cleaned_line)
            if opt_match:
                option_text = opt_match.group(2).strip()
                current_q["options"].append(option_text)
                continue

    if current_q:
        current_q["type"] = "tf" if len(current_q["options"]) == 2 else "mcq"
        questions.append(current_q)

    return questions

# =========================================================================
# 🚦 مسارات تسجيل الدخول والتسجيل
# =========================================================================
@app.route('/')
def index():
    if 'user_id' in session:
        role = session.get('role')
        if role == 'student': return redirect(url_for('student_dashboard'))
        elif role == 'parent': return redirect(url_for('parent_dashboard'))
        elif role == 'admin': return redirect(url_for('teacher_portal'))

    admin = users_col.find_one({"role": "admin"})
    teacher_avatar = admin.get('avatar') if admin and admin.get('avatar') else DEFAULT_AVATAR
    teacher_name = admin.get('name', 'الأستاذة هايدي عطية') if admin else 'الأستاذة هايدي عطية'

    return render_template('index.html', 
                           grades=GRADE_NAMES, 
                           teacher_avatar=teacher_avatar, 
                           teacher_name=teacher_name)

@app.route('/api/auth/register', methods=['POST'])
def register():
    data = request.json
    name = data.get('name', '').strip()
    phone = data.get('phone', '').strip()
    password = data.get('password', '').strip()
    role = data.get('role', 'student')
    grade = data.get('grade', 'p4')
    parent_phone = data.get('parent_phone', '').strip()

    if not name or not phone or not password:
        return jsonify({'status': 'error', 'msg': 'جميع الحقول مطلوبة!'}), 400

    if users_col.find_one({"phone": phone}):
        return jsonify({'status': 'error', 'msg': 'رقم الهاتف مسجل بالفعل!'}), 400

    student_code = f"HAIDY-{grade.upper()}-{int(datetime.now().timestamp()) % 10000}" if role == 'student' else None
    status = "pending" if role == 'student' else "approved"

    doc = {
        "name": name,
        "phone": phone,
        "password": generate_password_hash(password),
        "role": role,
        "grade": grade,
        "grade_name": GRADE_NAMES.get(grade, ''),
        "student_code": student_code,
        "parent_phone": parent_phone,
        "status": status,
        "created_at": datetime.now(timezone.utc)
    }

    res = users_col.insert_one(doc)

    if role == 'student':
        return jsonify({
            'status': 'pending',
            'msg': '✅ تم إرسال طلب انضمامك بنجاح! حسابك قيد المراجعة والاعتماد من الأستاذة هايدي عطية وسيتم تفعيله قريباً.'
        })

    session['user_id'] = str(res.inserted_id)
    session['name'] = name
    session['role'] = role
    session['grade'] = grade
    session['student_code'] = student_code

    return jsonify({'status': 'success', 'redirect': url_for('parent_dashboard')})

@app.route('/api/auth/login', methods=['POST'])
def login():
    data = request.json
    phone = data.get('phone', '').strip()
    password = data.get('password', '').strip()

    user = users_col.find_one({"phone": phone})
    if user and check_password_hash(user['password'], password):
        if user.get('role') == 'student':
            user_status = user.get('status', 'approved')
            if user_status == 'pending':
                return jsonify({
                    'status': 'error', 
                    'msg': '⏳ حسابك ما زال قيد المراجعة والموافقة من قِبل الأستاذة هايدي عطية. يرجى الانتظار حتى اعتماده.'
                }), 403
            elif user_status == 'rejected':
                return jsonify({
                    'status': 'error', 
                    'msg': '❌ تم رفض طلب الانضمام. يرجى مراجعة الأستاذة هايدي عطية.'
                }), 403

        session['user_id'] = str(user['_id'])
        session['name'] = user['name']
        session['role'] = user['role']
        session['grade'] = user.get('grade', 'p4')
        session['student_code'] = user.get('student_code', '')

        if user['role'] == 'student': redirect_url = url_for('student_dashboard')
        elif user['role'] == 'parent': redirect_url = url_for('parent_dashboard')
        elif user['role'] == 'admin': redirect_url = url_for('teacher_portal')
        else: redirect_url = url_for('index')

        return jsonify({'status': 'success', 'redirect': redirect_url})

    return jsonify({'status': 'error', 'msg': 'رقم الهاتف أو كلمة المرور غير صحيحة!'}), 401

@app.route('/logout')
def logout():
    session.clear()
    return redirect(url_for('index'))

# =========================================================================
# 🔄 مسارات فحص التحديثات التلقائية المباشرة (Polling Live Sync)
# =========================================================================
@app.route('/api/admin/live-check')
@login_required('admin')
def admin_live_check():
    grade = request.args.get('grade', 'p4')
    today = datetime.now().strftime("%Y-%m-%d")
    current_session = request.args.get('session', 'الحصة الأولى').strip()

    pending_count = users_col.count_documents({"grade": grade, "role": "student", "status": "pending"})
    students_count = users_col.count_documents({"grade": grade, "role": "student", "status": {"$ne": "pending"}})
    exams_count = exams_col.count_documents({"grade": grade})
    sessions_count = sessions_col.count_documents({"grade": grade, "date": today})
    attendance_count = attendance_col.count_documents({"grade": grade, "date": today, "session_name": current_session})
    
    # حساب إجمالي عدد تسليمات الامتحانات لهذا الصف
    grade_exams = list(exams_col.find({"grade": grade}, {"_id": 1}))
    exam_ids = [str(e["_id"]) for e in grade_exams]
    results_count = results_col.count_documents({"exam_id": {"$in": exam_ids}})

    signature = f"{pending_count}_{students_count}_{exams_count}_{sessions_count}_{attendance_count}_{results_count}"

    return jsonify({
        "status": "success",
        "signature": signature
    })

@app.route('/api/student/live-check')
@login_required('student')
def student_live_check():
    student_id = session.get('user_id')
    student_grade = session.get('grade', 'p4')

    exams_count = exams_col.count_documents({"grade": student_grade, "status": "published"})
    attendance_count = attendance_col.count_documents({"student_id": student_id})
    results_count = results_col.count_documents({"student_id": student_id})

    signature = f"{exams_count}_{attendance_count}_{results_count}"

    return jsonify({
        "status": "success",
        "signature": signature
    })

# =========================================================================
# 👨‍🎓 مسارات الطالب
# =========================================================================
@app.route('/student')
@login_required('student')
def student_dashboard():
    student_grade = session.get('grade', 'p4')
    student_id = session.get('user_id')
    now = datetime.now()

    active_exams = list(exams_col.find({"grade": student_grade, "status": "published"}).sort("created_at", -1))

    for exam in active_exams:
        exam['_id'] = str(exam['_id'])
        result = results_col.find_one({"exam_id": exam['_id'], "student_id": student_id})
        exam['taken'] = bool(result)
        exam['student_score'] = result['score'] if result else None
        
        if exam.get('deadline'):
            try:
                deadline_dt = datetime.strptime(exam['deadline'], "%Y-%m-%dT%H:%M")
                exam['is_expired'] = now > deadline_dt
            except Exception:
                exam['is_expired'] = False
        else:
            exam['is_expired'] = False

    my_results = list(results_col.find({"student_id": student_id}).sort("date", -1))
    my_attendance = list(attendance_col.find({"student_id": student_id}).sort("date", -1))

    initial_signature = f"{len(active_exams)}_{len(my_attendance)}_{len(my_results)}"

    return render_template('student_dashboard.html',
                           exams=active_exams,
                           results=my_results,
                           attendance=my_attendance,
                           grade_title=GRADE_NAMES.get(student_grade, ''),
                           initial_signature=initial_signature)

@app.route('/exam/<exam_id>')
@login_required('student')
def exam_room(exam_id):
    student_id = session.get('user_id')
    student_grade = session.get('grade', 'p4')

    exam = exams_col.find_one({"_id": ObjectId(exam_id), "grade": student_grade, "status": "published"})
    if not exam:
        return "الامتحان غير متاح لمرحلتك أو تم حذفه!", 404

    if results_col.find_one({"exam_id": str(exam['_id']), "student_id": student_id}):
        return redirect(url_for('student_dashboard'))

    if exam.get('deadline'):
        try:
            deadline_dt = datetime.strptime(exam['deadline'], "%Y-%m-%dT%H:%M")
            if datetime.now() > deadline_dt:
                return "عذراً، انتهى موعد تقديم هذا الامتحان وتم إغلاقه رسمياً.", 403
        except Exception:
            pass

    exam['_id'] = str(exam['_id'])
    return render_template('exam_room.html', exam=exam)

@app.route('/api/exam/submit', methods=['POST'])
@login_required('student')
def submit_student_exam():
    data = request.json
    exam_id = data.get('exam_id')
    answers = data.get('answers', {})
    cheat_report = data.get('cheat_report')

    student_id = session.get('user_id')
    student_name = session.get('name')
    student_grade = session.get('grade')

    exam = exams_col.find_one({"_id": ObjectId(exam_id)})
    if not exam:
        return jsonify({'status': 'error', 'msg': 'الامتحان غير موجود'}), 404

    if results_col.find_one({"exam_id": exam_id, "student_id": student_id}):
        return jsonify({'status': 'error', 'msg': 'تم تسليم هذا الامتحان مسبقاً!'}), 400

    questions = exam.get('questions', [])
    score = 0.0

    pts_mcq = float(exam.get('points_mcq', exam.get('points_per_q', 5)))
    pts_tf = float(exam.get('points_tf', pts_mcq))

    total_marks = 0.0
    for q in questions:
        q_type = q.get('type') or ('tf' if len(q.get('options', [])) == 2 else 'mcq')
        current_q_pts = pts_tf if q_type == 'tf' else pts_mcq
        total_marks += current_q_pts

        qid = str(q['id'])
        if answers.get(qid) == q.get('correct'):
            score += current_q_pts

    score_display = int(score) if score.is_integer() else round(score, 2)
    total_marks_display = int(total_marks) if total_marks.is_integer() else round(total_marks, 2)
    final_score = "محضر غش" if cheat_report else score_display

    results_col.insert_one({
        "exam_id": exam_id,
        "exam_title": exam['title'],
        "student_id": student_id,
        "student_name": student_name,
        "grade": student_grade,
        "grade_name": GRADE_NAMES.get(student_grade, ''),
        "score": final_score,
        "total_marks": total_marks_display,
        "cheat_report": cheat_report,
        "answers": answers,
        "date": datetime.now().strftime("%Y-%m-%d %H:%M")
    })

    return jsonify({'status': 'success', 'score': final_score, 'total': total_marks_display})

# =========================================================================
# 👨‍👩‍👧 مسارات ولي الأمر
# =========================================================================
@app.route('/parent')
@login_required('parent')
def parent_dashboard():
    parent_user = users_col.find_one({"_id": ObjectId(session['user_id'])})
    children = list(users_col.find({"parent_phone": parent_user['phone'], "role": "student"}))

    reports = []
    for ch in children:
        ch_id = str(ch['_id'])
        attendance = list(attendance_col.find({"student_id": ch_id}).sort("date", -1))
        results = list(results_col.find({"student_id": ch_id}).sort("date", -1))
        reports.append({
            "child": ch,
            "grade_title": GRADE_NAMES.get(ch.get('grade'), ''),
            "attendance": attendance,
            "results": results
        })

    return render_template('parent_dashboard.html', reports=reports)

# =========================================================================
# 👑 مسارات الأستاذة هايدي عطية
# =========================================================================
@app.route('/teacher-portal')
def teacher_portal():
    if session.get('role') != 'admin':
        return render_template('index.html', admin_prompt=True, grades=GRADE_NAMES)

    selected_grade = request.args.get('grade', 'p4')
    today = datetime.now().strftime("%Y-%m-%d")

    teacher = users_col.find_one({"role": "admin"}) or {}
    teacher_avatar = teacher.get('avatar') or DEFAULT_AVATAR
    teacher_name = teacher.get('name', 'الأستاذة هايدي عطية')

    pending_students = list(users_col.find({
        "role": "student", 
        "grade": selected_grade, 
        "status": "pending"
    }).sort("created_at", -1))
    for p in pending_students:
        p['_id'] = str(p['_id'])

    total_pending_all = users_col.count_documents({"role": "student", "status": "pending"})

    students = list(users_col.find({
        "role": "student", 
        "grade": selected_grade, 
        "status": {"$ne": "pending"}
    }).sort("name", 1))

    for s in students:
        s_id_str = str(s['_id'])
        s['_id'] = s_id_str
        
        att_history = list(attendance_col.find({"student_id": s_id_str}))
        p_count = sum(1 for a in att_history if a.get('status') == 'حاضر')
        a_count = sum(1 for a in att_history if a.get('status') == 'غائب')
        total = p_count + a_count
        s['present_count'] = p_count
        s['absent_count'] = a_count
        s['attendance_rate'] = round((p_count / total * 100), 1) if total > 0 else 100.0

    sessions_col.update_one(
        {"grade": selected_grade, "date": today, "name": "الحصة الأولى"},
        {"$setOnInsert": {"created_at": datetime.now(timezone.utc)}},
        upsert=True
    )

    session_docs = list(sessions_col.find({"grade": selected_grade, "date": today}).sort("created_at", 1))
    existing_sessions = [s['name'] for s in session_docs]

    req_session = request.args.get('session', '').strip()
    if req_session and req_session in existing_sessions:
        current_session = req_session
    else:
        current_session = existing_sessions[0] if existing_sessions else "الحصة الأولى"

    today_records = list(attendance_col.find({
        "date": today,
        "grade": selected_grade,
        "session_name": current_session
    }))
    today_status = {r['student_id']: r['status'] for r in today_records}

    exams = list(exams_col.find({"grade": selected_grade}).sort("created_at", -1))
    total_submissions = 0
    for ex in exams:
        ex_id_str = str(ex['_id'])
        ex['_id'] = ex_id_str
        ex_results = list(results_col.find({"exam_id": ex_id_str}).sort("date", -1))
        ex['submissions_count'] = len(ex_results)
        ex['results'] = ex_results
        total_submissions += len(ex_results)

    initial_signature = f"{len(pending_students)}_{len(students)}_{len(exams)}_{len(existing_sessions)}_{len(today_records)}_{total_submissions}"

    return render_template('admin_dashboard.html',
                           grades=GRADE_NAMES,
                           current_grade=selected_grade,
                           current_grade_title=GRADE_NAMES.get(selected_grade),
                           teacher_avatar=teacher_avatar,
                           teacher_name=teacher_name,
                           pending_students=pending_students,
                           total_pending_all=total_pending_all,
                           students=students,
                           today_status=today_status,
                           today=today,
                           existing_sessions=existing_sessions,
                           current_session=current_session,
                           exams=exams,
                           initial_signature=initial_signature)

@app.route('/api/admin/update-profile', methods=['POST'])
@login_required('admin')
def update_profile():
    data = request.json
    avatar = data.get('avatar')
    name = data.get('name', '').strip()

    update_fields = {}
    if avatar:
        update_fields['avatar'] = avatar
    if name:
        update_fields['name'] = name
        session['name'] = name

    if update_fields:
        users_col.update_one({"role": "admin"}, {"$set": update_fields})
        return jsonify({'status': 'success'})
    return jsonify({'status': 'error', 'msg': 'لم يتم إرسال أي تعديلات'}), 400

@app.route('/api/admin/approve-student/<student_id>', methods=['POST'])
@login_required('admin')
def approve_student(student_id):
    try:
        users_col.update_one(
            {"_id": ObjectId(student_id)},
            {"$set": {"status": "approved", "approved_at": datetime.now(timezone.utc)}}
        )
        return jsonify({'status': 'success'})
    except Exception as e:
        return jsonify({'status': 'error', 'msg': str(e)}), 400

@app.route('/api/admin/reject-student/<student_id>', methods=['POST'])
@login_required('admin')
def reject_student(student_id):
    try:
        users_col.delete_one({"_id": ObjectId(student_id)})
        return jsonify({'status': 'success'})
    except Exception as e:
        return jsonify({'status': 'error', 'msg': str(e)}), 400

@app.route('/api/admin/delete-student/<student_id>', methods=['POST'])
@login_required('admin')
def delete_student(student_id):
    try:
        users_col.delete_one({"_id": ObjectId(student_id)})
        attendance_col.delete_many({"student_id": student_id})
        results_col.delete_many({"student_id": student_id})
        return jsonify({'status': 'success'})
    except Exception as e:
        return jsonify({'status': 'error', 'msg': str(e)}), 400

@app.route('/api/admin/create-session', methods=['POST'])
@login_required('admin')
def create_session():
    data = request.json
    grade = data.get('grade')
    session_name = data.get('session_name', '').strip()
    today = datetime.now().strftime("%Y-%m-%d")

    if not session_name:
        return jsonify({'status': 'error', 'msg': 'اسم الحصة مطلوب'}), 400

    sessions_col.update_one(
        {"grade": grade, "date": today, "name": session_name},
        {"$setOnInsert": {"created_at": datetime.now(timezone.utc)}},
        upsert=True
    )
    return jsonify({'status': 'success'})

@app.route('/api/admin/delete-session', methods=['POST'])
@login_required('admin')
def delete_session():
    data = request.json
    grade = data.get('grade')
    session_name = data.get('session_name', '').strip()
    today = datetime.now().strftime("%Y-%m-%d")

    if not session_name or not grade:
        return jsonify({'status': 'error', 'msg': 'بيانات الحصة ناقصة'}), 400

    try:
        sessions_col.delete_one({"grade": grade, "date": today, "name": session_name})
        attendance_col.delete_many({"grade": grade, "date": today, "session_name": session_name})

        remaining = sessions_col.count_documents({"grade": grade, "date": today})
        if remaining == 0:
            sessions_col.insert_one({
                "grade": grade,
                "date": today,
                "name": "الحصة الأولى",
                "created_at": datetime.now(timezone.utc)
            })

        return jsonify({'status': 'success'})
    except Exception as e:
        return jsonify({'status': 'error', 'msg': str(e)}), 500

@app.route('/api/admin/attendance', methods=['POST'])
@login_required('admin')
def set_attendance():
    data = request.json
    student_id = data.get('student_id')
    status = data.get('status')
    grade = data.get('grade')
    session_name = data.get('session_name', 'الحصة الأولى').strip()
    today = datetime.now().strftime("%Y-%m-%d")

    student = users_col.find_one({"_id": ObjectId(student_id)})
    if not student:
        return jsonify({'status': 'error', 'msg': 'الطالب غير موجود'}), 404

    attendance_col.update_one(
        {"student_id": student_id, "date": today, "session_name": session_name},
        {"$set": {
            "student_name": student['name'],
            "grade": grade,
            "session_name": session_name,
            "status": status,
            "date": today,
            "updated_at": datetime.now(timezone.utc)
        }},
        upsert=True
    )
    return jsonify({'status': 'success'})

@app.route('/api/admin/parse-word-exam', methods=['POST'])
@login_required('admin')
def parse_word_exam():
    if 'file' not in request.files:
        return jsonify({'status': 'error', 'msg': 'يرجى اختيار ملف Word (.docx)'}), 400

    file = request.files['file']
    if not file.filename.lower().endswith('.docx'):
        return jsonify({'status': 'error', 'msg': 'الملف يجب أن يكون بصيغة docx'}), 400

    text_content = extract_text_from_docx(file)
    questions = parse_mcq_content(text_content)

    if not questions:
        return jsonify({'status': 'error', 'msg': 'لم يتم العثور على أسئلة بتنسيق واضح في الملف'}), 400

    return jsonify({
        'status': 'success',
        'questions_count': len(questions),
        'questions': questions
    })

@app.route('/api/admin/publish-exam', methods=['POST'])
@login_required('admin')
def publish_exam():
    data = request.json
    title = data.get('title')
    grade = data.get('grade')
    duration = int(data.get('duration', 20))
    deadline = data.get('deadline')
    questions = data.get('questions', [])
    points_mcq = float(data.get('points_mcq', 5))
    points_tf = float(data.get('points_tf', 2.5))

    if not title or not questions:
        return jsonify({'status': 'error', 'msg': 'بيانات الامتحان أو الأسئلة ناقصة'}), 400

    total_marks = sum(
        points_tf if (q.get('type') == 'tf' or len(q.get('options', [])) == 2) else points_mcq 
        for q in questions
    )

    doc = {
        "title": title,
        "grade": grade,
        "grade_name": GRADE_NAMES.get(grade),
        "duration": duration,
        "deadline": deadline,
        "questions": questions,
        "points_mcq": points_mcq,
        "points_tf": points_tf,
        "total_marks": int(total_marks) if total_marks.is_integer() else round(total_marks, 2),
        "status": "published",
        "created_at": datetime.now(timezone.utc)
    }
    exams_col.insert_one(doc)
    return jsonify({'status': 'success'})

@app.route('/api/admin/delete-exam/<exam_id>', methods=['POST'])
@login_required('admin')
def delete_exam(exam_id):
    try:
        exams_col.delete_one({"_id": ObjectId(exam_id)})
        results_col.delete_many({"exam_id": exam_id})
        return jsonify({'status': 'success'})
    except Exception as e:
        return jsonify({'status': 'error', 'msg': str(e)}), 400

@app.route('/api/admin/get-exam/<exam_id>')
@login_required('admin')
def get_exam_details(exam_id):
    try:
        exam = exams_col.find_one({"_id": ObjectId(exam_id)})
        if not exam:
            return jsonify({'status': 'error', 'msg': 'الامتحان غير موجود'}), 404
        exam['_id'] = str(exam['_id'])
        if 'created_at' in exam: exam['created_at'] = str(exam['created_at'])
        if 'updated_at' in exam: exam['updated_at'] = str(exam['updated_at'])
        return jsonify({'status': 'success', 'exam': exam})
    except Exception as e:
        return jsonify({'status': 'error', 'msg': str(e)}), 400

@app.route('/api/admin/update-exam/<exam_id>', methods=['POST'])
@login_required('admin')
def update_exam(exam_id):
    data = request.json
    title = data.get('title')
    duration = int(data.get('duration', 20))
    deadline = data.get('deadline')
    questions = data.get('questions', [])
    points_mcq = float(data.get('points_mcq', 5))
    points_tf = float(data.get('points_tf', 2.5))

    total_marks = sum(
        points_tf if (q.get('type') == 'tf' or len(q.get('options', [])) == 2) else points_mcq 
        for q in questions
    )

    try:
        exams_col.update_one(
            {"_id": ObjectId(exam_id)},
            {"$set": {
                "title": title,
                "duration": duration,
                "deadline": deadline,
                "questions": questions,
                "points_mcq": points_mcq,
                "points_tf": points_tf,
                "total_marks": int(total_marks) if total_marks.is_integer() else round(total_marks, 2),
                "updated_at": datetime.now(timezone.utc)
            }}
        )
        return jsonify({'status': 'success'})
    except Exception as e:
        return jsonify({'status': 'error', 'msg': str(e)}), 400

if __name__ == '__main__':
    app.run(debug=True, port=5000)
